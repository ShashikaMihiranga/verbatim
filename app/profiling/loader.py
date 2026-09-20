"""Read a spreadsheet file into a clean object-dtype pandas DataFrame.

Pure functions, no network. Everything the loader changes is written to
`LoadedTable.notes` so the UI can show it (nothing is done silently).
Cells are stripped strings, native numbers/dates (xlsx) or None for missing.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

MAX_ROWS = 500_000
MAX_COLS = 300
MAX_UNCOMPRESSED_BYTES = 400 * 1024 * 1024
MAX_ZIP_RATIO = 200

MISSING_TOKENS = {
    "", "n/a", "na", "n.a.", "nan", "null", "nil", "-", "--", "?", "#n/a", "#null!",
    "not applicable", "missing",
}
DELIMS = [",", ";", "\t", "|"]

_NUM_ONLY = re.compile(r"^[-+]?[$€£¥]?\s*\d[\d,.]*\s*%?$")


class LoadError(ValueError):
    """A file we cannot read. The message is safe to show to the user."""


@dataclass
class LoadedTable:
    df: pd.DataFrame
    notes: list[str] = field(default_factory=list)
    meta: dict[str, Any] = field(default_factory=dict)


def _clean_cell(v: Any, missing_counter: Counter) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        s = " ".join(v.replace(" ", " ").split())
        if s.lower() in MISSING_TOKENS:
            if s:
                missing_counter[s] += 1
            return None
        return s
    if isinstance(v, float) and v != v:
        return None
    if isinstance(v, datetime) and v.time() == datetime.min.time():
        return v.date()
    return v


def _is_numeric_like(v: Any) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    return isinstance(v, str) and bool(_NUM_ONLY.match(v))


def _header_like(row: list[Any]) -> bool:
    cells = [c for c in row if c is not None]
    if not cells:
        return False
    texty = sum(1 for c in cells if isinstance(c, str) and not _is_numeric_like(c))
    return texty / len(cells) >= 0.6 and len({str(c).lower() for c in cells}) >= 0.8 * len(cells)


def _rows_to_table(raw_rows: list[list[Any]], notes: list[str]) -> pd.DataFrame:
    missing: Counter = Counter()
    kept: list[list[Any]] = []
    blank = 0
    for raw in raw_rows:
        r = [_clean_cell(c, missing) for c in raw]
        if all(c is None for c in r):
            blank += 1
        else:
            kept.append(r)
    if not kept:
        raise LoadError("The file has no data in it.")

    width = max(len(r) for r in kept)
    if width > MAX_COLS:
        raise LoadError(f"The file has more than {MAX_COLS} columns, which is too many to check.")
    kept = [r + [None] * (width - len(r)) for r in kept]

    # header row detection: first row that is mostly full and mostly words
    probe = kept[:15]
    fill = [sum(c is not None for c in r) for r in probe]
    top = max(fill)
    header_idx = None
    for i, r in enumerate(probe):
        if fill[i] >= 0.6 * top and _header_like(r):
            header_idx = i
            break
    if header_idx is None:
        header: list[Any] = [None] * width
        body = kept
        notes.append("No header row found, so columns are named Column 1, Column 2 and so on.")
    else:
        header = list(kept[header_idx])
        body = kept[header_idx + 1:]
        if header_idx:
            notes.append(f"Skipped {header_idx} row(s) above the header row.")

    cols = list(zip(*body)) if body else [() for _ in range(width)]
    keep_idx = [i for i, c in enumerate(cols) if any(v is not None for v in c)]
    dropped_empty = width - len(keep_idx)

    names, seen, unnamed = [], Counter(), 0
    for i in keep_idx:
        h = header[i]
        name = str(h).strip() if h is not None else ""
        if not name:
            name = f"Column {i + 1}"
            unnamed += 1
        seen[name] += 1
        if seen[name] > 1:
            name = f"{name} ({seen[name]})"
        names.append(name)

    if header_idx is not None and unnamed:
        notes.append("Some columns had no header (or a merged header), so they were named Column N.")
    if dropped_empty:
        notes.append(f"Dropped {dropped_empty} completely empty column(s).")
    if blank:
        notes.append(f"Ignored {blank} blank row(s).")
    if missing:
        total = sum(missing.values())
        toks = ", ".join(f'"{t}"' for t, _ in missing.most_common(4))
        notes.append(f"Treated {total} cell(s) reading {toks} as blank.")

    if not keep_idx:
        raise LoadError("The file has a header but no answers under it.")
    return pd.DataFrame({n: pd.Series(list(cols[i]), dtype=object) for n, i in zip(names, keep_idx)})


def _decode(data: bytes, notes: list[str]) -> tuple[str, str]:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16"), "utf-16"
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8"), "utf-8-sig"
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        notes.append("The file is not UTF-8; read as Windows-1252 (Excel's usual save format).")
        return data.decode("cp1252", errors="replace"), "cp1252"


def _sniff_delimiter(text: str) -> str:
    lines = [ln for ln in text.splitlines()[:40] if ln.strip()]
    best, best_score = ",", (0.0, 0)
    for d in DELIMS:
        counts = [len(r) for r in csv.reader(lines, delimiter=d)]
        if not counts:
            continue
        mode, freq = Counter(counts).most_common(1)[0]
        if mode < 2:
            continue
        score = (freq / len(counts), mode)
        if score > best_score:
            best, best_score = d, score
    return best


def load_csv(path: Path) -> LoadedTable:
    notes: list[str] = []
    data = path.read_bytes()
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise LoadError("This does not look like a CSV file (it contains binary data).")
    text, enc = _decode(data, notes)
    delim = _sniff_delimiter(text)
    csv.field_size_limit(1_000_000)
    raw: list[list[str]] = []
    try:
        for i, r in enumerate(csv.reader(io.StringIO(text, newline=""), delimiter=delim)):
            if i >= MAX_ROWS:
                raise LoadError(f"The file has more than {MAX_ROWS:,} rows, which is too many.")
            raw.append(r)
    except csv.Error as e:
        raise LoadError(f"Could not read the CSV: {e}") from e
    if delim != ",":
        notes.append(f"Columns are separated by {'a tab' if delim == chr(9) else repr(delim)}.")
    df = _rows_to_table(raw, notes)
    return LoadedTable(df, notes, {"format": "csv", "encoding": enc, "delimiter": delim})


def _check_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            total = sum(i.file_size for i in infos)
            comp = sum(i.compress_size for i in infos) or 1
            if total > MAX_UNCOMPRESSED_BYTES or (total > 50 * 1024 * 1024 and total / comp > MAX_ZIP_RATIO):
                raise LoadError("This Excel file expands to an unsafe size, so it was refused.")
    except zipfile.BadZipFile as e:
        raise LoadError("This is not a valid .xlsx file.") from e


def load_xlsx(path: Path) -> LoadedTable:
    import openpyxl

    _check_zip(path)
    notes: list[str] = []
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise LoadError("Could not open this .xlsx file.") from e
    try:
        chosen, raw = None, []
        for ws in wb.worksheets:
            rows = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS:
                    raise LoadError(f"The sheet has more than {MAX_ROWS:,} rows, which is too many.")
                rows.append(list(r))
            if any(c is not None for r in rows for c in r):
                chosen, raw = ws.title, rows
                break
        if chosen is None:
            raise LoadError("The workbook has no data in it.")
        if len(wb.worksheets) > 1:
            notes.append(
                f'Read the sheet "{chosen}" ({len(wb.worksheets)} sheets in the file; only the first with data is used).'
            )
    finally:
        wb.close()
    df = _rows_to_table(raw, notes)
    return LoadedTable(df, notes, {"format": "xlsx", "sheet": chosen})


def load_xls(path: Path) -> LoadedTable:
    try:
        import xlrd
    except ImportError as e:  # pragma: no cover
        raise LoadError("Old .xls files need the xlrd package installed.") from e
    notes: list[str] = []
    try:
        wb = xlrd.open_workbook(str(path))
    except Exception as e:
        raise LoadError("This is not a valid .xls file.") from e
    sheet = next((s for s in wb.sheets() if s.nrows), None)
    if sheet is None:
        raise LoadError("The workbook has no data in it.")
    if sheet.nrows > MAX_ROWS:
        raise LoadError(f"The sheet has more than {MAX_ROWS:,} rows, which is too many.")
    raw = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            v = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    v = xlrd.xldate.xldate_as_datetime(v, wb.datemode)
                except Exception:
                    pass
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                v = None
            elif cell.ctype == xlrd.XL_CELL_NUMBER and float(v).is_integer():
                v = int(v)
            row.append(v)
        raw.append(row)
    df = _rows_to_table(raw, notes)
    return LoadedTable(df, notes, {"format": "xls", "sheet": sheet.name})


def load_table(path: str | Path) -> LoadedTable:
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".csv", ".tsv", ".txt"):
        return load_csv(path)
    if ext == ".xlsx":
        return load_xlsx(path)
    if ext == ".xls":
        return load_xls(path)
    raise LoadError("Only .csv, .xlsx and .xls files are supported.")
